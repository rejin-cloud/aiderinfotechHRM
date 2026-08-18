import calendar
from datetime import date, datetime
import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.attendance.models import Attendance


def generate_monthly_attendance_excel(users, year, month, department=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    month_name = calendar.month_name[month]
    ws.title = f"Attendance {month_name[:3]} {year}"
    ws.views.sheetView[0].showGridLines = True

    # Styling Palettes
    title_font = Font(name='Segoe UI', size=16, bold=True, color='FFFFFF')
    title_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')

    meta_font = Font(name='Segoe UI', size=10, italic=True, color='475569')
    header_font = Font(name='Segoe UI', size=9, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')

    cell_font = Font(name='Segoe UI', size=9)
    bold_cell_font = Font(name='Segoe UI', size=9, bold=True)
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Status fills for daily matrix
    status_fills = {
        'PRESENT': PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid'),  # Light green
        'LATE': PatternFill(start_color='FEF9C3', end_color='FEF9C3', fill_type='solid'),     # Light yellow
        'HALF_DAY': PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid'), # Light blue
        'ON_LEAVE': PatternFill(start_color='EDE9FE', end_color='EDE9FE', fill_type='solid'), # Light purple
        'ABSENT': PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid'),   # Light red
    }

    # Title Banner
    ws.merge_cells('A1:AJ1')
    ws['A1'] = f"Aider Infotech • Monthly Attendance Report ({month_name} {year})"
    ws['A1'].font = title_font
    ws['A1'].fill = title_fill
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[1].height = 36

    dept_label = department.name if department else "All Corporate Departments"
    ws['A2'] = f"Organization Scope: {dept_label} | Generated on: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    ws['A2'].font = meta_font
    ws.row_dimensions[2].height = 20

    # Days in month
    num_days = calendar.monthrange(year, month)[1]

    # Header Row
    headers = ["Emp ID", "Employee Name", "Department", "Role"]
    for d in range(1, num_days + 1):
        headers.append(f"{d:02d}")
    headers.extend(["Present", "Late", "Half-Day", "On Leave", "Absent", "Total Hrs"])

    start_row = 4
    ws.row_dimensions[start_row].height = 24
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align if col_num > 4 else left_align
        cell.border = thin_border

    # Fetch attendance records for this month and users
    attendances = Attendance.objects.filter(
        user__in=users,
        date__year=year,
        date__month=month
    ).select_related('user')

    # Build lookup map: (user_id, day) -> attendance_record
    att_map = {}
    for att in attendances:
        att_map[(att.user_id, att.date.day)] = att

    current_row = start_row + 1
    for u in users:
        ws.row_dimensions[current_row].height = 20
        # Basic user info
        ws.cell(row=current_row, column=1, value=u.employee_id or f"EMP-{u.id}").alignment = left_align
        ws.cell(row=current_row, column=2, value=u.get_full_name() or u.username).alignment = left_align
        ws.cell(row=current_row, column=3, value=u.department.name if u.department else "Unassigned").alignment = left_align
        ws.cell(row=current_row, column=4, value=u.get_role_display()).alignment = left_align

        for c in range(1, 5):
            cell = ws.cell(row=current_row, column=c)
            cell.font = cell_font
            cell.border = thin_border

        # Days columns
        present_cnt = 0
        late_cnt = 0
        half_day_cnt = 0
        leave_cnt = 0
        absent_cnt = 0
        total_mins = 0

        for d in range(1, num_days + 1):
            col_idx = 4 + d
            att = att_map.get((u.id, d))
            cell = ws.cell(row=current_row, column=col_idx)
            cell.alignment = center_align
            cell.border = thin_border
            cell.font = cell_font

            if att:
                total_mins += att.work_duration_minutes
                if att.status in [Attendance.Status.PRESENT, Attendance.Status.LATE, Attendance.Status.HALF_DAY] or att.is_present:
                    present_cnt += 1
                if att.status == Attendance.Status.PRESENT:
                    cell.value = "P"
                    cell.fill = status_fills['PRESENT']
                elif att.status == Attendance.Status.LATE or att.is_late:
                    cell.value = "L"
                    cell.fill = status_fills['LATE']
                    late_cnt += 1
                elif att.status == Attendance.Status.HALF_DAY:
                    cell.value = "HD"
                    cell.fill = status_fills['HALF_DAY']
                    half_day_cnt += 1
                elif att.status == Attendance.Status.ON_LEAVE:
                    cell.value = "LV"
                    cell.fill = status_fills['ON_LEAVE']
                    leave_cnt += 1
                else:
                    cell.value = "A"
                    cell.fill = status_fills['ABSENT']
                    absent_cnt += 1
            else:
                # Check if it is a weekend
                day_of_week = calendar.weekday(year, month, d)
                if day_of_week in [5, 6]:  # Saturday, Sunday
                    cell.value = "—"
                else:
                    cell.value = "A"
                    cell.fill = status_fills['ABSENT']
                    absent_cnt += 1

        # Summary columns
        sum_col_start = 4 + num_days + 1
        summary_values = [
            present_cnt,
            late_cnt,
            half_day_cnt,
            leave_cnt,
            absent_cnt,
            round(total_mins / 60, 1)
        ]
        for idx, val in enumerate(summary_values):
            cell = ws.cell(row=current_row, column=sum_col_start + idx, value=val)
            cell.alignment = center_align
            cell.border = thin_border
            cell.font = bold_cell_font

        current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        col_idx = col[0].column
        if col_idx in [1, 2, 3, 4]:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        elif col_idx <= 4 + num_days:
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 4.5
        else:
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 10

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_monthly_attendance_pdf(users, year, month, department=None):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=20,
        leftMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A')
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#64748B')
    )
    th_style = ParagraphStyle(
        'Th',
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )
    td_style = ParagraphStyle(
        'Td',
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0F172A')
    )
    td_center = ParagraphStyle(
        'TdCenter',
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0F172A'),
        alignment=1
    )

    month_name = calendar.month_name[month]
    dept_name = department.name if department else "All Departments"

    story = [
        Paragraph(f"Aider Infotech • Monthly Attendance Ledger ({month_name} {year})", title_style),
        Paragraph(f"Department Scope: <b>{dept_name}</b> | Generated on: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", subtitle_style),
        HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563EB'), spaceAfter=10),
    ]

    # Attendance summary table
    table_data = [
        [
            Paragraph("<b>Emp ID</b>", th_style),
            Paragraph("<b>Employee Name</b>", th_style),
            Paragraph("<b>Department</b>", th_style),
            Paragraph("<b>Role</b>", th_style),
            Paragraph("<b>Present</b>", th_style),
            Paragraph("<b>Late</b>", th_style),
            Paragraph("<b>Half-Day</b>", th_style),
            Paragraph("<b>On Leave</b>", th_style),
            Paragraph("<b>Absent</b>", th_style),
            Paragraph("<b>Total Work (Hrs)</b>", th_style),
        ]
    ]

    attendances = Attendance.objects.filter(
        user__in=users,
        date__year=year,
        date__month=month
    )
    att_by_user = {}
    for att in attendances:
        att_by_user.setdefault(att.user_id, []).append(att)

    for u in users:
        u_records = att_by_user.get(u.id, [])
        p_cnt = sum(1 for a in u_records if a.status in [Attendance.Status.PRESENT, Attendance.Status.LATE, Attendance.Status.HALF_DAY] or a.is_present)
        l_cnt = sum(1 for a in u_records if a.status == Attendance.Status.LATE or a.is_late)
        hd_cnt = sum(1 for a in u_records if a.status == Attendance.Status.HALF_DAY)
        lv_cnt = sum(1 for a in u_records if a.status == Attendance.Status.ON_LEAVE)
        ab_cnt = sum(1 for a in u_records if a.status == Attendance.Status.ABSENT)
        tot_mins = sum(a.work_duration_minutes for a in u_records)

        table_data.append([
            Paragraph(u.employee_id or f"EMP-{u.id}", td_style),
            Paragraph(u.get_full_name() or u.username, td_style),
            Paragraph(u.department.name if u.department else "—", td_style),
            Paragraph(u.get_role_display(), td_style),
            Paragraph(str(p_cnt), td_center),
            Paragraph(str(l_cnt), td_center),
            Paragraph(str(hd_cnt), td_center),
            Paragraph(str(lv_cnt), td_center),
            Paragraph(str(ab_cnt), td_center),
            Paragraph(f"{round(tot_mins/60, 1)} hrs", td_center),
        ])

    table = Table(
        table_data,
        colWidths=[65, 140, 110, 100, 50, 45, 55, 55, 50, 80],
        repeatRows=1
    )
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))

    story.append(table)
    doc.build(story)
    buffer.seek(0)
    return buffer
